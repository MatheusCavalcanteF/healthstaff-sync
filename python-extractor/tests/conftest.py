from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
import pytest


@dataclass
class PaginaPdf:
    """Página sintética: linhas de texto no topo e, opcionalmente, uma tabela
    com bordas desenhadas (detectável pela estratégia 'lines' do pdfplumber)."""

    linhas: list[str] = field(default_factory=list)
    tabela: list[list[str]] | None = None
    larguras: list[float] | None = None
    paisagem: bool = False


def _escapar(texto: str) -> str:
    return texto.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _larguras_padrao(tabela: list[list[str]]) -> list[float]:
    n_colunas = max(len(linha) for linha in tabela)
    larguras = []
    for c in range(n_colunas):
        maior = max(len(linha[c]) if c < len(linha) else 0 for linha in tabela)
        larguras.append(max(16.0, maior * 3.6 + 6))
    return larguras


def _conteudo_pagina(pagina: PaginaPdf, altura: float) -> str:
    partes = []
    if pagina.linhas:
        texto = " T* ".join(f"({_escapar(linha)}) Tj" for linha in pagina.linhas)
        partes.append(f"BT /F1 10 Tf 12 TL 30 {altura - 30} Td {texto} ET")

    if pagina.tabela:
        larguras = pagina.larguras or _larguras_padrao(pagina.tabela)
        altura_linha = 12.0
        topo = altura - 50 - 12 * len(pagina.linhas)
        partes.append("0.5 w")
        for r, linha in enumerate(pagina.tabela):
            y = topo - (r + 1) * altura_linha
            x = 20.0
            for c, largura in enumerate(larguras):
                partes.append(f"{x:.1f} {y:.1f} {largura:.1f} {altura_linha:.1f} re S")
                celula = linha[c] if c < len(linha) else ""
                if celula:
                    partes.append(
                        f"BT /F1 6 Tf {x + 2:.1f} {y + 3.5:.1f} Td "
                        f"({_escapar(celula)}) Tj ET"
                    )
                x += largura
    return "\n".join(partes)


def gerar_pdf(paginas: list[list[str] | PaginaPdf]) -> bytes:
    """Gera um PDF mínimo e válido. Cada página é uma lista de linhas de
    texto ou um ``PaginaPdf`` (texto + tabela)."""
    paginas = [p if isinstance(p, PaginaPdf) else PaginaPdf(linhas=p) for p in paginas]
    objetos: list[bytes] = []
    n_paginas = len(paginas)
    fonte_id = 3 + 2 * n_paginas
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(n_paginas))

    objetos.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objetos.append(f"<< /Type /Pages /Kids [{kids}] /Count {n_paginas} >>".encode())
    for i, pagina in enumerate(paginas):
        largura, altura = (842, 595) if pagina.paisagem else (595, 842)
        objetos.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {largura} {altura}] "
            f"/Contents {4 + 2 * i} 0 R "
            f"/Resources << /Font << /F1 {fonte_id} 0 R >> >> >>".encode()
        )
        stream = _conteudo_pagina(pagina, altura).encode("latin-1")
        objetos.append(
            b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream"
        )
    objetos.append(
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica "
        b"/Encoding /WinAnsiEncoding >>"
    )

    saida = bytearray(b"%PDF-1.4\n")
    offsets = []
    for numero, corpo in enumerate(objetos, start=1):
        offsets.append(len(saida))
        saida += b"%d 0 obj\n" % numero + corpo + b"\nendobj\n"
    inicio_xref = len(saida)
    saida += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objetos) + 1)
    for offset in offsets:
        saida += b"%010d 00000 n \n" % offset
    saida += (
        b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n"
        % (len(objetos) + 1, inicio_xref)
    )
    return bytes(saida)


@pytest.fixture
def criar_pdf(tmp_path: Path) -> Callable[..., Path]:
    """Fábrica: criar_pdf(linhas | PaginaPdf, nome=..., paginas_extras=...) -> Path."""

    def _criar(
        primeira: list[str] | PaginaPdf,
        nome: str = "escala.pdf",
        paginas_extras: list[list[str] | PaginaPdf] | None = None,
    ) -> Path:
        caminho = tmp_path / nome
        caminho.write_bytes(gerar_pdf([primeira, *(paginas_extras or [])]))
        return caminho

    return _criar


@pytest.fixture
def criar_xlsx(tmp_path: Path) -> Callable[..., Path]:
    """Fábrica: criar_xlsx({aba: linhas}, nome=..., mesclar={aba: ["G8:AK8"]}) -> Path.

    Mesclar mantém só o valor da primeira célula, como no Excel real.
    """

    def _criar(
        abas: dict[str, list[list]],
        nome: str = "planilha.xlsx",
        mesclar: dict[str, list[str]] | None = None,
    ) -> Path:
        livro = openpyxl.Workbook()
        livro.remove(livro.active)
        for titulo, linhas in abas.items():
            planilha = livro.create_sheet(titulo)
            for linha in linhas:
                planilha.append(list(linha))
            for intervalo in (mesclar or {}).get(titulo, []):
                planilha.merge_cells(intervalo)
        caminho = tmp_path / nome
        livro.save(caminho)
        return caminho

    return _criar
